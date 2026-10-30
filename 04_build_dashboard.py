"""Step 4: Build the Excel dashboard (native Excel charts + live formulas + household dropdown)."""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
A = HERE / "analysis_outputs"
S = HERE / "sql_outputs"

monthly = pd.read_csv(S / "monthly_by_household.csv")
hourly = pd.read_csv(S / "hourly_profile.csv")
yearly = pd.read_csv(S / "yearly_summary.csv")
seasonal = pd.read_csv(S / "seasonal_summary.csv")
sens = pd.read_csv(A / "weather_sensitivity.csv")
adj = pd.read_csv(A / "weather_adjusted_yearly.csv")
dq_b = pd.read_csv(S / "dq_before.csv").iloc[0]
dq_a = pd.read_csv(S / "dq_after.csv").iloc[0]

FONT = "Arial"
NAVY, LIGHT, GREY = "1F4E79", "DCE6F1", "F2F2F2"
hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor=NAVY)
base = Font(name=FONT, size=10)
thin = Side(style="thin", color="BFBFBF")


def write_table(ws, df, start_row=1, start_col=1, num_formats=None):
    num_formats = num_formats or {}
    for j, col in enumerate(df.columns):
        c = ws.cell(start_row, start_col + j, col)
        c.font, c.fill = hdr_font, hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            if pd.isna(v):
                v = None
            c = ws.cell(start_row + i, start_col + j, v)
            c.font = base
            if df.columns[j] in num_formats:
                c.number_format = num_formats[df.columns[j]]
    for j, col in enumerate(df.columns):
        ws.column_dimensions[get_column_letter(start_col + j)].width = max(12, len(str(col)) + 3)
    ws.freeze_panes = ws.cell(start_row + 1, start_col)


wb = Workbook()
dash = wb.active
dash.title = "Dashboard"

# ------------------------------------------------------------------ data sheets
ws = wb.create_sheet("Monthly")
write_table(ws, monthly, num_formats={"total_kwh": "#,##0.0", "avg_hourly_kwh": "0.000"})

ws = wb.create_sheet("Hourly")
write_table(ws, hourly, num_formats={"avg_kwh": "0.000"})

ws = wb.create_sheet("Yearly")
write_table(ws, yearly, num_formats={"total_kwh": "#,##0", "yoy_pct_change": '0.0"%"'})

ws = wb.create_sheet("Seasonal")
write_table(ws, seasonal, num_formats={"avg_daily_kwh": "0.00"})

ws = wb.create_sheet("Weather")
write_table(ws, sens, num_formats={c: "0.000" for c in sens.columns[1:]})
ws.cell(len(sens) + 3, 1, "Model: daily kWh = base + a*HDD + b*CDD (HDD base 15C, CDD base 22C), fit by least squares in Pandas/NumPy.").font = Font(name=FONT, italic=True, size=9)
write_table(ws, adj, start_row=len(sens) + 6, num_formats={"weather_adjusted_daily_kwh": "0.00"})
ws.freeze_panes = None

ws = wb.create_sheet("Data Quality")
dq = pd.DataFrame(
    {
        "check": ["Total rows (raw)", "Duplicate rows", "Missing kWh", "Negative kWh", "Implausible spikes (>15 kWh/hr)",
                  "Missing temperature (kept; not needed for kWh)", "Rows in clean table", "Rows removed"],
        "count": [dq_b.total_rows, dq_b.duplicate_rows, dq_b.missing_kwh, dq_b.negative_kwh, dq_b.implausible_spikes,
                  dq_b.missing_temp, dq_a.clean_rows, dq_a.rows_removed],
    }
)
write_table(ws, dq, num_formats={"count": "#,##0"})
ws.column_dimensions["A"].width = 48
ws.cell(len(dq) + 3, 1, "Cleaning rules (SQL): de-duplicate on household+timestamp; drop NULL, negative and >15 kWh/hr readings.").font = Font(name=FONT, italic=True, size=9)

# ------------------------------------------------------------------ Dashboard
dash.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKLMNOP", [3, 16, 14, 14, 14, 14, 14, 14, 3, 14, 14, 14, 14, 14, 14, 3]):
    dash.column_dimensions[col].width = w

dash["B1"] = "Household Electricity Consumption Dashboard"
dash["B1"].font = Font(name=FONT, bold=True, size=18, color=NAVY)
dash["B2"] = "Hourly smart-meter readings, Jan 2023 - Dec 2025. Choose a household in the yellow cell; everything below updates."
dash["B2"].font = Font(name=FONT, italic=True, size=9, color="595959")

dash["B4"] = "Household:"
dash["B4"].font = Font(name=FONT, bold=True, size=11)
dash["C4"] = "H001"
dash["C4"].font = Font(name=FONT, bold=True, size=12, color="0000FF")
dash["C4"].fill = PatternFill("solid", fgColor="FFFF00")
dash["C4"].alignment = Alignment(horizontal="center")
dv = DataValidation(type="list", formula1='"H001,H002,H003"', allow_blank=False)
dash.add_data_validation(dv)
dv.add("C4")
dash["D4"] = '=IF(C4="H001","Detached, 4 occupants (EV from Jul 2024)",IF(C4="H002","Condo, 1 occupant",IF(C4="H003","Townhouse, 3 occupants","")))'
dash["D4"].font = Font(name=FONT, size=10, color="595959")

# --- chart-feed block (month list is in col B..; data pulled with SUMIFS so it follows the dropdown)
FEED = 47  # first feed row (below the charts)
dash.cell(FEED - 2, 2, "Chart data (feeds the charts above - formulas, do not overwrite)").font = Font(name=FONT, bold=True, size=9, color="7F7F7F")
months = sorted(monthly.year_month.unique())
# Month block: B, C  (cols 2,3)
for col_i, h in [(2, "Month"), (3, "kWh")]:
    c = dash.cell(FEED - 1, col_i, h); c.font, c.fill = hdr_font, hdr_fill
for i, m in enumerate(months):
    r = FEED + i
    dash.cell(r, 2, m).font = base
    c = dash.cell(r, 3, f"=SUMIFS(Monthly!$D:$D,Monthly!$A:$A,$C$4,Monthly!$C:$C,B{r})")
    c.font, c.number_format = base, "#,##0"
# Hour block: E,F,G (cols 5,6,7)
for col_i, h in [(5, "Hour"), (6, "Weekday"), (7, "Weekend")]:
    c = dash.cell(FEED - 1, col_i, h); c.font, c.fill = hdr_font, hdr_fill
for hr in range(24):
    r = FEED + hr
    dash.cell(r, 5, hr).font = base
    for col_i, dt in [(6, "Weekday"), (7, "Weekend")]:
        c = dash.cell(r, col_i, f'=SUMIFS(Hourly!$D:$D,Hourly!$A:$A,$C$4,Hourly!$B:$B,"{dt}",Hourly!$C:$C,$E{r})')
        c.font, c.number_format = base, "0.00"
# Year block: J,K (cols 10,11)
for col_i, h in [(10, "Year"), (11, "Annual kWh")]:
    c = dash.cell(FEED - 1, col_i, h); c.font, c.fill = hdr_font, hdr_fill
for i, y in enumerate([2023, 2024, 2025]):
    r = FEED + i
    dash.cell(r, 10, str(y)).font = base
    c = dash.cell(r, 11, f"=SUMIFS(Yearly!$C:$C,Yearly!$A:$A,$C$4,Yearly!$B:$B,{y})")
    c.font, c.number_format = base, "#,##0"
# Season block: M,N (cols 13,14)
for col_i, h in [(13, "Season"), (14, "Avg kWh/day")]:
    c = dash.cell(FEED - 1, col_i, h); c.font, c.fill = hdr_font, hdr_fill
for i, s in enumerate(["1 Winter", "2 Spring", "3 Summer", "4 Fall"]):
    r = FEED + i
    dash.cell(r, 13, s[2:]).font = base
    c = dash.cell(r, 14, f'=SUMIFS(Seasonal!$C:$C,Seasonal!$A:$A,$C$4,Seasonal!$B:$B,"{s}")')
    c.font, c.number_format = base, "0.0"

m_first, m_last = FEED, FEED + len(months) - 1

# --- KPI tiles (formulas)
kpis = [
    ("B", "Total kWh (3 yrs)", f"=SUM(C{m_first}:C{m_last})", "#,##0"),
    ("D", "Avg kWh / month", f"=AVERAGE(C{m_first}:C{m_last})", "#,##0"),
    ("F", "Peak month kWh", f"=MAX(C{m_first}:C{m_last})", "#,##0"),
    ("H", "2025 vs 2023", f"=K{FEED+2}/K{FEED}-1", "+0.0%;-0.0%"),
]
for col, label, formula, fmt in kpis:
    ci = ord(col) - 64
    dash.merge_cells(start_row=6, start_column=ci, end_row=6, end_column=ci + 1)
    dash.merge_cells(start_row=7, start_column=ci, end_row=7, end_column=ci + 1)
    l = dash.cell(6, ci, label)
    l.font = Font(name=FONT, size=9, color="595959"); l.alignment = Alignment(horizontal="center")
    v = dash.cell(7, ci, formula)
    v.font = Font(name=FONT, bold=True, size=18, color=NAVY); v.number_format = fmt
    v.alignment = Alignment(horizontal="center")
    for r in (6, 7):
        for cc in (ci, ci + 1):
            dash.cell(r, cc).fill = PatternFill("solid", fgColor=LIGHT)
dash.row_dimensions[7].height = 30
dash["H7"].value = f"=K{FEED+2}/K{FEED}-1"

# peak-hour KPI
dash["J6"] = "Weekday peak hour"
dash["J6"].font = Font(name=FONT, size=9, color="595959"); dash["J6"].alignment = Alignment(horizontal="center")
dash["J7"] = f"=INDEX(E{FEED}:E{FEED+23},MATCH(MAX(F{FEED}:F{FEED+23}),F{FEED}:F{FEED+23},0))"
dash["J7"].number_format = '0":00"'
dash["J7"].font = Font(name=FONT, bold=True, size=18, color=NAVY); dash["J7"].alignment = Alignment(horizontal="center")
dash["J6"].fill = dash["J7"].fill = PatternFill("solid", fgColor=LIGHT)

# --- charts
def style(ch, title, w=17, h=7.5):
    ch.title = title; ch.width, ch.height = w, h
    ch.legend.position = "b"

lc = LineChart(); style(lc, "Monthly consumption (kWh)")
lc.add_data(Reference(dash, min_col=3, min_row=FEED - 1, max_row=m_last), titles_from_data=True)
lc.set_categories(Reference(dash, min_col=2, min_row=m_first, max_row=m_last))
lc.legend = None; lc.series[0].smooth = False; lc.y_axis.title = "kWh"; lc.x_axis.tickLblSkip = 3
lc.series[0].graphicalProperties.line.solidFill = NAVY; lc.series[0].graphicalProperties.line.width = 28000
lc.x_axis.delete = False; lc.y_axis.delete = False
dash.add_chart(lc, "B9")

hc = LineChart(); style(hc, "Average load by hour of day (kWh)")
hc.add_data(Reference(dash, min_col=6, max_col=7, min_row=FEED - 1, max_row=FEED + 23), titles_from_data=True)
hc.set_categories(Reference(dash, min_col=5, min_row=FEED, max_row=FEED + 23))
hc.series[0].graphicalProperties.line.solidFill = NAVY
hc.series[1].graphicalProperties.line.solidFill = "C0504D"
hc.series[1].graphicalProperties.line.dashStyle = "dash"
hc.series[0].smooth = False; hc.series[1].smooth = False; hc.x_axis.title = "Hour"; hc.x_axis.delete = False; hc.y_axis.delete = False
dash.add_chart(hc, "J9")

bc = BarChart(); style(bc, "Annual consumption (kWh)", w=17, h=7.5)
bc.add_data(Reference(dash, min_col=11, min_row=FEED - 1, max_row=FEED + 2), titles_from_data=True)
bc.set_categories(Reference(dash, min_col=10, min_row=FEED, max_row=FEED + 2))
bc.legend = None; bc.series[0].graphicalProperties.solidFill = NAVY
bc.x_axis.delete = False; bc.y_axis.delete = False
dash.add_chart(bc, "B25")

sc = BarChart(); style(sc, "Avg daily kWh by season", w=17, h=7.5)
sc.add_data(Reference(dash, min_col=14, min_row=FEED - 1, max_row=FEED + 3), titles_from_data=True)
sc.set_categories(Reference(dash, min_col=13, min_row=FEED, max_row=FEED + 3))
sc.legend = None; sc.series[0].graphicalProperties.solidFill = "C0504D"
sc.x_axis.delete = False; sc.y_axis.delete = False
dash.add_chart(sc, "J25")

dash["B42"] = "Source: SQLite energy.db (readings_clean) via Python/Pandas. Data is synthetic, built to mimic smart-meter behaviour."
dash["B42"].font = Font(name=FONT, italic=True, size=8, color="7F7F7F")

# ------------------------------------------------------------------ Key findings sheet
ws = wb.create_sheet("Findings", 1)
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 110
ws["B1"] = "Key findings"; ws["B1"].font = Font(name=FONT, bold=True, size=16, color=NAVY)
notes = [
    "Daily shape: H002 and H003 peak at 19:00 on weekdays and weekends; evening is the heaviest period for typical homes.",
    "Weather drives usage: winter days use roughly 2x the electricity of summer days; temperature explains 80-99% of daily variation (R-squared on the Weather sheet).",
    "H001 jumped ~18% in 2024 and ~14% in 2025: overnight (23:00-03:00) load rose by ~1.8 kWh/hr starting July 2024, consistent with EV charging.",
    "After removing weather effects, H001's underlying use rose ~56% (2023 to 2025), while H002 fell ~8.5% (appliance/lighting efficiency) and H003 was flat.",
    "Data quality: 687 of 79,062 raw rows (0.9%) were removed - 150 duplicates, 473 missing, 39 negative and 25 spike (unit-error) readings.",
    "Recommendation angle: EV charging makes 23:00 H001's weekday peak hour; smart-scheduling the charge across later overnight hours would flatten that peak.",
]
for i, n in enumerate(notes, start=3):
    c = ws.cell(i, 2, f"{i-2}.  {n}")
    c.font = Font(name=FONT, size=11); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[i].height = 34

# header styling for data sheets
for name in ["Monthly", "Hourly", "Yearly", "Seasonal", "Weather", "Data Quality"]:
    wb[name].row_dimensions[1].height = 22

dash.page_setup.orientation = "landscape"; dash.page_setup.fitToWidth = 1; dash.page_setup.fitToHeight = 0
dash.sheet_properties.pageSetUpPr.fitToPage = True
out = HERE / "Energy_Dashboard.xlsx"
wb.save(out)
print("saved", out)
